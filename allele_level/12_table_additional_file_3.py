import csv,collections,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
R=list(csv.DictReader(open('uni/coding_alleles_validated.tsv'),delimiter='\t'))
G=list(csv.DictReader(open('uni/final_gene_classes.tsv'),delimiter='\t'))
wb=openpyxl.Workbook()
def sheet(ws,header,rows):
    ws.append(header)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.freeze_panes='A2'
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(45,max(10,max(len(str(c.value or ''))*0.9 for c in col[:200])))
ws=wb.active; ws.title='README'
for t in ["Additional file 3. Allele-level comparison and read-level validation of HIGH- and MODERATE-impact coding alleles in the combined set of 1,370 candidate genes (GHB-732, LR; GHB-538, HR).",
 "Combined candidate set: union of the final candidate lists (323 common, 262 GHB-732-only, 785 GHB-538-only). Variants within each gene model (mRNA ± 100 bp) were taken from both varieties' VCFs and annotated with SnpEff 5.3 against the same gene models.",
 "Genotype QC (column 'reliable'): DP >= 10, GQ >= 20; heterozygous allele fraction 0.25-0.75 or homozygous-alternate >= 0.85; not in a paralog-like cluster (>= 4 heterozygous coding calls with allele fraction < 0.35 within 100 bp in the same gene).",
 "Read-level check (column 'read_check'), on reads with MAPQ >= 20 (secondary, supplementary and duplicate-flagged reads excluded) at each reliable allele not called identically in both varieties:",
 "  specific = >= 15% alternate-like support in own reads; other variety >= 10 informative reads with < 5% alternate-like and < 25% total non-reference support; neither alignment > 2x its median depth at inspected sites (41 reads GHB-732; 68 reads GHB-538).",
 "  present_in_other_reads = >= 15% alternate-like support in the other variety; excess_depth = > 2x median depth in either alignment; not_supported_by_own_reads, other_uncovered (< 10 reads), other_divergent_site (>= 25% other non-reference support) and ambiguous (5-15%) = unresolved.",
 "  Alternate-like support: alternate base (SNV/MNV); insertion/deletion of similar length within 10 bp; for indels >= 15 bp also soft-clipping within 10 bp. self_/other_ columns give informative read counts and fractions.",
 "Gene classes: 'Complex locus' = any paralog-flagged coding call, or median coding-variant DP > 2x the variety median (46 GHB-732; 70 GHB-538) in either variety. Otherwise classified from validated specific alleles (LR-specific, HR-specific, Both varieties), else 'Shared alleles only' or 'No reliable coding difference'.",
 "Functional categories are those of Additional file 4 (one rule set for both varieties)."]: ws.append([t])
ws.column_dimensions['A'].width=180
keys=['variety','gene','chrom','pos','ref','alt','impact','effect','hgvs_c','hgvs_p','GT','DP','GQ','AF','hom','paralog_flag','reliable','status_in_other','read_check','self_reads','self_alt_frac','other_reads','other_alt_frac','other_nonref_frac']
sheet(wb.create_sheet('Coding alleles'),keys,[[r.get(k,'') for k in keys] for r in R])
gk=list(G[0].keys())
sheet(wb.create_sheet('Gene classes'),gk,[[g[k] for k in gk] for g in G])
cls=['LR-specific','HR-specific','Both varieties','Shared alleles only','No reliable coding difference','Complex locus']
T=collections.Counter((g['functional_category'],g['allele_class']) for g in G)
cats=sorted(set(g['functional_category'] for g in G),key=lambda c:(c.startswith('Not'),c))
sheet(wb.create_sheet('Category x class'),['Functional category']+cls+['Total'],[[c]+[T[(c,k)] for k in cls]+[sum(T[(c,k)] for k in cls)] for c in cats]+[['Total']+[sum(T[(c,k)] for c in cats) for k in cls]+[len(G)]])
CH=collections.Counter((g['chrom'],g['allele_class']) for g in G)
sheet(wb.create_sheet('Chromosome x class'),['Chromosome']+cls+['Total'],[[c]+[CH[(c,k)] for k in cls]+[sum(CH[(c,k)] for k in cls)] for c in [f'chr{i}' for i in range(1,8)]])
wb.save('out/Additional_file_3.xlsx'); print('ok')
