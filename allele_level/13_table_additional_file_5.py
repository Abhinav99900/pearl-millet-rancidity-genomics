import json,csv,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
rows=json.load(open('uni/ranking2.json'))
G={r['gene_id']:r for r in csv.DictReader(open('uni/final_gene_classes.tsv'),delimiter='\t')}
wb=openpyxl.Workbook()
def sheet(ws,header,data):
    ws.append(header); [ws.append(r) for r in data]
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.freeze_panes='A2'
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(50,max(10,max(len(str(c.value or '')) for c in col)*0.9))
ws=wb.active; ws.title='README'
for t in ['Additional file 5. Priority-scoring output and class-wise ranking for validation-panel selection (allele-level version).',
 'Priority score = 2 x impact score + 2 x rancidity-relevance score (Table 6). Impact score = that of the most severe validated variety-specific coding allele of the locus (HIGH = 4, MODERATE = 3); see Additional file 3 for allele-level validation.',
 'Eligibility: >= 1 validated variety-specific coding allele; not a complex locus (no paralog-like cluster and median coding-variant depth <= 2x the variety median in both varieties); predicted protein >= 70% of the length of the best Swiss-Prot hit (UniProtKB, accessed October 2026).',
 'Within each class the highest-scoring eligible locus was selected; ties: (i) more validated variety-specific coding alleles; (ii) best hit with documented seed-storage-lipid function; (iii) higher allele fraction of the scoring allele in own reads.',
 'Sheets: Ranking (all 184 loci); Tie-breaking; PgTAGLip1-3 status; Primers (Table 8).']: ws.append([t])
ws.column_dimensions['A'].width=170
hdr=['family','gene_id','best Swiss-Prot hit','protein_aa','best-hit length (aa)','model coverage','allele-level class','validated specific alleles (n)','GHB-732-specific (n)','GHB-538-specific (n)','scoring allele','scoring allele variety','impact','genotype','own-read alt fraction','other-variety reads','other-variety alt fraction','relevance','priority score','eligibility','rank among eligible','selected']
order={'TAG lipase':0,'MAG lipase':1,'Lipoxygenase':2,'Class III peroxidase':3,'Superoxide dismutase':4}
rows.sort(key=lambda r:(order.get(r['family'],9),r.get('rank_eligible',999),-(r.get('score') or 0),r['gene_id']))
sheet(wb.create_sheet('Ranking'),hdr,[[r['family'],r['gene_id'],r['homolog'],r['protein_aa'],r['hit_aa'],r['model_coverage'],r['allele_class'],r['n_validated_specific'],r['n_LR'],r['n_HR'],r.get('best_allele'),r.get('best_allele_variety'),r.get('impact'),r.get('GT'),r.get('own_read_alt_frac'),r.get('other_reads'),r.get('other_alt_frac'),r['relevance'],r.get('score'),r['eligibility'],r.get('rank_eligible'),r.get('selected','')] for r in rows])
sheet(wb.create_sheet('Tie-breaking'),['Class','Tied loci (score)','Criterion applied','Values','Selected'],[
 ['TAG lipase','Pgl_GLEAN_10017888 (14) - only eligible locus','none','-','Pgl_GLEAN_10017888 (SDP1)'],
 ['MAG lipase','Pgl_GLEAN_10036216 (14); Pgl_GLEAN_10013892 (14)','(i) equal (1 vs 1); (ii) neither; (iii) own-read allele fraction','0.482 vs 0.243','Pgl_GLEAN_10036216'],
 ['Class III peroxidase','Pgl_GLEAN_10033116 (14); Pgl_GLEAN_10004812 (14)','(i) number of validated specific alleles','2 vs 1','Pgl_GLEAN_10033116 (PER52-like)'],
 ['Lipoxygenase','no eligible locus','-','3 complex (10005615, 10010905, 10028100); 2 truncated (10032393, 112/870 aa; 10003609, 606/899 aa)','none'],
 ['Superoxide dismutase','no eligible locus','-','1 truncated (10001023, 84/151 aa)','none']])
tg=[('PgTAGLip1','Pgl_GLEAN_10023115','No validated variety-specific coding allele; single GHB-538 missense call p.Gln282Arg absent from GHB-732 reads but at a site with excess GHB-732 depth (86 reads; > 2x median)'),
    ('PgTAGLip2','Pgl_GLEAN_10024115','Complex locus (1 paralog-flagged call; 1.89x depth in GHB-732); p.Ala223Val (GHB-732) at an excess-depth site'),
    ('PgTAGLip3','Pgl_GLEAN_10011213','Complex locus: 46 paralog-flagged coding calls, 2.28x depth in GHB-732; HIGH-impact calls (stop/frameshift) at similar low allele fractions in both varieties; two GHB-732-specific missense alleles (p.Arg347Trp, p.His332Gln) not interpreted')]
sheet(wb.create_sheet('PgTAGLip1-3 status'),['Gene (Aher et al. [10])','Locus','Chromosome','In GHB-732 list','In GHB-538 list','Allele-level class','Paralog-flagged coding calls','Depth ratio GHB-732','Depth ratio GHB-538','Allele-level interpretation'],
 [[n,g,G[g]['chrom'],G[g]['in_732_list'],G[g]['in_538_list'],G[g]['allele_class'],G[g]['paralog_flagged_variants'],G[g]['depth_ratio_732'],G[g]['depth_ratio_538'],t] for n,g,t in tg])
sheet(wb.create_sheet('Primers (Table 8)'),['Gene','Locus','Forward (5\'-3\')','Fwd Tm (°C)','Fwd GC (%)','Reverse (5\'-3\')','Rev Tm (°C)','Rev GC (%)','Amplicon (bp)','CDS position','Spans exon-exon junction','Primer-BLAST (C. americanus, nt)','Min. mismatches to nearest non-target candidate CDS (F / R)','Variants under primers (either variety)','Note'],[
 ['SDP1','Pgl_GLEAN_10017888','GCAGGTACATGAGGCGGATT',60.18,55,'CAGAACACCACATCCCGGAG',60.39,60,105,'66-170','No','No off-target product','5 / 5','None','No variant in amplicon'],
 ['MAGL','Pgl_GLEAN_10036216','GGCAGGCATCTTTGTACCCT',60.03,55,'AGAATATCGGTGCCACAGCC',60.18,55,117,'395-511','Yes','No off-target product','5 / 4','None','No variant in amplicon'],
 ['PER52-like POX','Pgl_GLEAN_10033116','CGGTGCATCCAGTACAAGGA',59.75,55,'ACTTCCAAAGGCATCCCTGG',59.96,55,243,'391-633','Yes','No off-target product','4 / 4','None','Product > 150 bp because of tandem paralog Pgl_GLEAN_10033115; amplicon includes GHB-732 p.Ser180fs site and other variant positions']])
wb.save('out/Additional_file_5.xlsx'); print('ok')
